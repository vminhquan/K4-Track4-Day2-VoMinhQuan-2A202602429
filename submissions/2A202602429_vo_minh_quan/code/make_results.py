"""Bước 5: gom mọi số liệu từ log chạy thật -> results.xlsx (7 sheet theo GUIDE 6.1) + biểu đồ tổng hợp
+ results/tables.md (bảng markdown để dán vào report.md).

Chạy từ thư mục bài nộp:  python code/make_results.py
Nguồn: runs/*/seed*/summary.json, results/backbones.csv, results/inference_val.csv, results/latency.csv,
       results/eval/* (do eval.py sinh), results/selection.json, results/final.json.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
from dataset import CLASS_NAMES  # noqa: E402
from experiments import BASELINE_SEEDS, TRAINING  # noqa: E402

R = Path("results")
F = Path("figures")


def summ(exp, seed=0):
    p = Path("runs") / exp / f"seed{seed}" / "summary.json"
    return json.loads(p.read_text()) if p.exists() else None


def pm(mean, std, d=4):
    return f"{mean:.{d}f} ± {std:.{d}f}"


def sheet_backbones():
    df = pd.read_csv(R / "backbones.csv")
    df = df.rename(columns={
        "weight_tag": "tag trọng số (timm)", "params_M": "#tham số (M)", "gmacs": "GMAC (fvcore)",
        "img_size": "độ phân giải", "val_macro_f1": "macro-F1 val", "val_top1": "top-1 val",
        "train_time_per_epoch_s": "thời gian train/epoch (s)", "lat_b1_p50_ms": "độ trễ b1 p50 (ms)",
        "lat_b1_p95_ms": "độ trễ b1 p95 (ms)", "lat_b1_p99_ms": "độ trễ b1 p99 (ms)",
        "thr_b32_imgs_s": "thông lượng b32 (ảnh/s)", "device": "thiết bị đo", "curve": "ảnh biểu đồ"})
    sel = json.loads((R / "selection.json").read_text())
    df["ghi chú"] = ["ĐƯỢC CHỌN cho Bước 2: " + sel["backbone_reason"] if b == sel["backbone"] else "1 seed"
                     for b in df["backbone"]]
    # biểu đồ F1 theo độ trễ
    fig, ax = plt.subplots(figsize=(7.5, 4.8))
    ax.scatter(df["độ trễ b1 p50 (ms)"], df["macro-F1 val"], s=df["#tham số (M)"] * 8)
    for _, r in df.iterrows():
        ax.annotate(f"{r.exp_id} {r.backbone}", (r["độ trễ b1 p50 (ms)"], r["macro-F1 val"]), fontsize=8,
                    xytext=(4, 4), textcoords="offset points")
    ax.set(xlabel=f"độ trễ batch 1 p50 (ms) · {df['thiết bị đo'].iloc[0]}", ylabel="macro-F1 val",
           title="Bước 1: macro-F1 val theo độ trễ (kích thước điểm ∝ số tham số)")
    ax.grid(alpha=0.3)
    fig.tight_layout()
    fig.savefig(F / "backbones_f1_vs_latency.png", dpi=120)
    plt.close(fig)
    return df


def sheet_training():
    sel = json.loads((R / "selection.json").read_text())
    base = summ("T00", 0)["val_macro_f1"]
    t00 = [summ("T00", s) for s in BASELINE_SEEDS if summ("T00", s)]
    std = float(np.std([s["val_macro_f1"] for s in t00], ddof=1)) if len(t00) > 1 else float("nan")
    rows = []
    meta = {e: (ax, diff) for e, ax, diff, _ in TRAINING}
    for c in sel.get("combos", []):
        meta[c["exp_id"]] = ("kết hợp", c["desc"])
    for eid, (axis, diff) in meta.items():
        seeds = BASELINE_SEEDS if eid == "T00" else (0,)
        for s in seeds:
            x = summ(eid, s)
            if x is None:
                continue
            d = x["val_macro_f1"] - base
            rows.append({
                "exp_id": eid, "backbone": x["weight_tag"], "trục": axis, "khác T00 ở điểm nào": diff, "seed": s,
                "macro-F1 val": x["val_macro_f1"], "top-1 val": x["val_top1"], "balanced acc val": x["val_bal_acc"],
                "Δ macro-F1 vs T00 seed0": d if eid != "T00" else 0.0 if s == 0 else d,
                "|Δ| / std(T00)": abs(d) / std if std == std and std > 0 else np.nan,
                "F1 Chinee apple val": x["val_f1_per_class"][0], "F1 Snake weed val": x["val_f1_per_class"][7],
                "F1 Negatives val": x["val_f1_per_class"][8], "best epoch": x["best_epoch"],
                "thời gian train/epoch (s)": x["train_time_per_epoch_s"], "ảnh biểu đồ": x["curve"],
                "ghi chú": ("mốc; 3 seed để đo nhiễu" if eid == "T00" else
                            "vượt nhiễu (Δ > std T00)" if d > std else
                            "kém hơn rõ (Δ < -std T00)" if d < -std else "trong khoảng nhiễu: không phân biệt được")
                + (" · ĐƯỢC CHỌN làm công thức chung kết" if eid == sel.get("recipe_exp") and s == 0 else ""),
            })
    df = pd.DataFrame(rows)
    f1s = [s["val_macro_f1"] for s in t00]
    agg = {"exp_id": "T00 (mean ± std)", "seed": f"{len(f1s)} seed", "macro-F1 val": pm(np.mean(f1s), std),
           "top-1 val": pm(np.mean([s["val_top1"] for s in t00]), np.std([s["val_top1"] for s in t00], ddof=1))}
    df = pd.concat([df, pd.DataFrame([agg])], ignore_index=True)
    # biểu đồ ablation
    d1 = df[(df.seed == 0) & (df.exp_id != "T00")].copy()
    fig, ax = plt.subplots(figsize=(10, 4.8))
    ax.bar(d1.exp_id, d1["Δ macro-F1 vs T00 seed0"], color=["tab:green" if v > std else "tab:red" if v < -std
                                                              else "tab:gray" for v in d1["Δ macro-F1 vs T00 seed0"]])
    ax.axhspan(-std, std, color="orange", alpha=0.2, label=f"±1 std T00 qua 3 seed ({std:.4f})")
    ax.axhline(0, color="black", lw=0.8)
    ax.set(ylabel="Δ macro-F1 val so với T00 (seed 0)", title=f"Bước 2: ablation công thức huấn luyện ({sel['backbone']})")
    ax.legend()
    ax.grid(axis="y", alpha=0.3)
    fig.tight_layout()
    fig.savefig(F / "training_ablation.png", dpi=120)
    plt.close(fig)
    return df, std


def sheet_inference():
    df = pd.read_csv(R / "inference_val.csv")
    return df.rename(columns={"macro_f1": "macro-F1 val", "top1": "top-1 val", "ece": "ECE val",
                              "p50_ms": "p50 b1 (ms)", "p95_ms": "p95 b1 (ms)", "p99_ms": "p99 b1 (ms)",
                              "imgs_per_s_b1": "thông lượng b1 (ảnh/s)", "cost_vs_I00": "chi phí tương đối (p50/I00)",
                              "delta_f1_vs_I00": "Δ macro-F1 vs I00", "checkpoint": "mô hình/checkpoint"})


def sheet_latency():
    df = pd.read_csv(R / "latency.csv")
    return df.rename(columns={"gpu": "thiết bị", "fused_bn": "gộp BN", "p50": "p50 (ms)", "p95": "p95 (ms)",
                              "p99": "p99 (ms)", "images_per_s": "ảnh/s", "includes_preprocessing": "tính tiền xử lý"})


def read_eval(tag):
    p = R / "eval" / f"{tag}_summary.json"
    return json.loads(p.read_text()) if p.exists() else None


def sheet_final():
    final = json.loads((R / "final.json").read_text())
    sel = json.loads((R / "selection.json").read_text())
    rows = []
    desc = {
        "F01": f"{sel['backbone']} + {sel['recipe_exp']} {sel['recipe_overrides']} + {sel['inference']} + TS",
        "F01_uncal": f"như F01, không temperature scaling",
        "F01rt": f"{sel['backbone']} + {sel['recipe_exp']} + 1 view + TS (thời gian thực)",
        "T00": f"mốc: {sel['backbone']} + công thức nền + 1 view (I00)",
    }
    val_of = {}
    for s in final["per_seed"]:
        val_of[s["seed"]] = s["val_macro_f1"]
    t00_val = {s: summ("T00", s)["val_macro_f1"] for s in BASELINE_SEEDS}
    for tag in ("F01", "F01_uncal", "F01rt", "T00"):
        p = R / "eval" / f"{tag}_per_seed.csv"
        if not p.exists():
            continue
        per = pd.read_csv(p)
        for _, r in per.iterrows():
            v = val_of.get(r.seed) if tag == "F01" else t00_val.get(r.seed) if tag == "T00" else None
            rows.append({"exp_id": tag, "cấu hình": desc[tag], "seed": int(r.seed), "macro-F1 val": v,
                         "macro-F1 test": r.macro_f1, "top-1 test": r.top1, "balanced acc test": r.balanced_acc,
                         "ECE test": r.ece, "NLL test": r.nll})
        e = read_eval(tag)
        vals = [x["macro-F1 val"] for x in rows if x["exp_id"] == tag and x["macro-F1 val"] is not None]
        rows.append({"exp_id": f"{tag} (mean ± std, {len(per)} seed)", "cấu hình": desc[tag], "seed": "tổng hợp",
                     "macro-F1 val": pm(np.mean(vals), np.std(vals, ddof=1)) if len(vals) > 1 else None,
                     "macro-F1 test": pm(e["macro_f1"]["mean"], e["macro_f1"]["std"]),
                     "top-1 test": pm(e["top1"]["mean"], e["top1"]["std"]),
                     "balanced acc test": pm(e["balanced_acc"]["mean"], e["balanced_acc"]["std"]),
                     "ECE test": pm(e["ece"]["mean"], e["ece"]["std"]), "NLL test": pm(e["nll"]["mean"], e["nll"]["std"])})
    return pd.DataFrame(rows)


def sheet_perclass():
    rows = []
    for tag, label in (("F01", "chung kết F01"), ("T00", "mốc T00+I00"), ("F01rt", "thời gian thực F01rt")):
        p = R / "eval" / f"{tag}_per_class.csv"
        if not p.exists():
            continue
        for _, r in pd.read_csv(p).iterrows():
            rows.append({"cấu hình": label, "lớp": r["class"], "số ảnh test": int(r.support),
                         "precision": pm(r.precision_mean, r.precision_std, 3),
                         "recall": pm(r.recall_mean, r.recall_std, 3), "F1": pm(r.f1_mean, r.f1_std, 3),
                         "recall (mean)": r.recall_mean, "F1 (mean)": r.f1_mean})
    return pd.DataFrame(rows)


def sheet_summary(bdf, tdf, idf):
    rows = []
    for _, r in bdf.iterrows():
        rows.append({"exp_id": r.exp_id, "loại": "backbone", "mô tả": r["tag trọng số (timm)"],
                     "macro-F1 val": r["macro-F1 val"], "top-1 val": r["top-1 val"],
                     "độ trễ b1 p95 (ms)": r["độ trễ b1 p95 (ms)"], "chi phí": f"{r['GMAC (fvcore)']:.2f} GMAC"})
    for _, r in tdf[tdf.seed.isin([0])].iterrows():
        rows.append({"exp_id": r.exp_id, "loại": "huấn luyện", "mô tả": r["khác T00 ở điểm nào"],
                     "macro-F1 val": r["macro-F1 val"], "top-1 val": r["top-1 val"], "độ trễ b1 p95 (ms)": None,
                     "chi phí": f"{r['thời gian train/epoch (s)']:.0f} s/epoch"})
    for _, r in idf.dropna(subset=["macro-F1 val"]).iterrows():
        if str(r.note).startswith("thành viên"):
            continue
        rows.append({"exp_id": r.exp_id, "loại": "suy luận", "mô tả": r.method, "macro-F1 val": r["macro-F1 val"],
                     "top-1 val": r["top-1 val"], "độ trễ b1 p95 (ms)": r.get("p95 b1 (ms)"),
                     "chi phí": f"K={r.K}"})
    df = pd.DataFrame(rows).sort_values("macro-F1 val", ascending=False).head(10).reset_index(drop=True)
    df.insert(0, "hạng", range(1, len(df) + 1))
    return df


def style(writer, name, df, highlight_col=None, best="max"):
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    ws = writer.sheets[name]
    ws.freeze_panes = "B2"
    for c in ws[1]:
        c.font = Font(bold=True)
    for i, col in enumerate(df.columns, 1):
        width = max(len(str(col)), *(len(str(v)) for v in df[col].head(50))) if len(df) else len(str(col))
        ws.column_dimensions[get_column_letter(i)].width = min(60, width + 2)
        if pd.api.types.is_float_dtype(df[col]):
            for cell in ws[get_column_letter(i)][1:]:
                cell.number_format = "0.0000"
    if highlight_col and highlight_col in df and len(df):
        vals = pd.to_numeric(df[highlight_col], errors="coerce")
        if vals.notna().any():
            idx = vals.idxmax() if best == "max" else vals.idxmin()
            for c in ws[idx + 2]:
                c.fill = PatternFill("solid", fgColor="FFF2A8")


def main():
    F.mkdir(exist_ok=True)
    bdf = sheet_backbones()
    tdf, std = sheet_training()
    idf = sheet_inference()
    ldf = sheet_latency()
    fdf = sheet_final()
    pdf = sheet_perclass()
    sdf = sheet_summary(bdf, tdf, idf)
    # khối so sánh chung kết với mốc dưới bảng Summary
    e_f, e_b = read_eval("F01"), read_eval("T00")
    comp = pd.DataFrame([
        {"hạng": "", "exp_id": "TEST", "loại": "chung kết F01", "mô tả": "mean ± std 3 seed",
         "macro-F1 val": pm(e_f["macro_f1"]["mean"], e_f["macro_f1"]["std"]),
         "top-1 val": pm(e_f["top1"]["mean"], e_f["top1"]["std"]), "chi phí": "cột là chỉ số TEST"},
        {"hạng": "", "exp_id": "TEST", "loại": "mốc T00 + I00", "mô tả": "mean ± std 3 seed",
         "macro-F1 val": pm(e_b["macro_f1"]["mean"], e_b["macro_f1"]["std"]),
         "top-1 val": pm(e_b["top1"]["mean"], e_b["top1"]["std"]), "chi phí": "cột là chỉ số TEST"},
        {"hạng": "", "exp_id": "TEST", "loại": "Δ (F01 - T00)", "mô tả": "",
         "macro-F1 val": f"{e_f['macro_f1']['mean'] - e_b['macro_f1']['mean']:+.4f}",
         "top-1 val": f"{e_f['top1']['mean'] - e_b['top1']['mean']:+.4f}", "chi phí": ""},
    ]) if e_f and e_b else pd.DataFrame()

    with pd.ExcelWriter("results.xlsx", engine="openpyxl") as w:
        for name, df, hl in [("Summary", sdf, "macro-F1 val"), ("Backbones", bdf, "macro-F1 val"),
                             ("Training", tdf, "macro-F1 val"), ("Inference", idf, "macro-F1 val"),
                             ("Final", fdf, "macro-F1 test"), ("PerClass", pdf, None), ("Latency", ldf, None)]:
            df.to_excel(w, sheet_name=name, index=False)
            style(w, name, df, hl)
        if len(comp):
            comp.to_excel(w, sheet_name="Summary", index=False, startrow=len(sdf) + 3)
            ws = w.sheets["Summary"]
            ws.cell(row=len(sdf) + 3, column=1, value="So sánh chung kết với mốc trên TEST (eval.py)")

    with open(R / "tables.md", "w") as f:
        for name, df in [("Backbones", bdf), ("Training", tdf), ("Inference", idf), ("Final", fdf),
                         ("PerClass", pdf), ("Latency", ldf), ("Summary", sdf)]:
            f.write(f"\n## {name}\n\n{df.to_markdown(index=False, floatfmt='.4f')}\n")
    print("Đã ghi results.xlsx và results/tables.md")
    # kiểm tra mỗi thí nghiệm huấn luyện có ảnh biểu đồ
    missing = [c for c in list(bdf["ảnh biểu đồ"]) + list(tdf["ảnh biểu đồ"].dropna()) if not (Path("curves") / c).exists()]
    print("Thiếu ảnh biểu đồ:", missing or "không")


if __name__ == "__main__":
    main()
